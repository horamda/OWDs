"""OWDs 2026 · Del Palacio — web de seguimiento de OWDs.

Lee el Google Sheet de OWDs público o con una cuenta de servicio y
sirve el tablero en static/index.html. El Sheet es la única fuente de
datos: respuestas de los formularios, hoja PERSONAS, programación y
REPROGRAMACION.
"""
import json
import os
import threading
import time
from functools import wraps
from pathlib import Path
from datetime import date, datetime
from io import BytesIO
from urllib.error import HTTPError, URLError
from urllib.request import urlopen

from openpyxl import load_workbook

from flask import Flask, Response, jsonify, request, send_from_directory
from google.oauth2 import service_account
from googleapiclient.discovery import build

SHEET_ID = os.environ.get("SHEET_ID", "1q5Q54vzSTUAuRnxVjtg87Bjohn0NwK2p2rrkzNZvld0")
CACHE_SECONDS = int(os.environ.get("CACHE_SECONDS", "60"))
SHEET_ACCESS = os.environ.get("SHEET_ACCESS", "public").strip().lower()
ACCESS_PASSWORD = os.environ.get("ACCESS_PASSWORD", "")
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]

app = Flask(__name__, static_folder="static", static_url_path="/static")

_cache = {"at": 0.0, "data": None}
_lock = threading.Lock()
_service = None


def sheets_service():
    """Cliente de Sheets con credenciales del entorno o del archivo local."""
    global _service
    if _service is None:
        raw = os.environ.get("GOOGLE_CREDENTIALS_JSON")
        if not raw:
            credentials_path = Path(__file__).resolve().parent / "credenciales.json"
            if not credentials_path.is_file():
                raise RuntimeError(
                    "Faltan credenciales: configurá GOOGLE_CREDENTIALS_JSON "
                    "o agregá credenciales.json a la carpeta del proyecto."
                )
            raw = credentials_path.read_text(encoding="utf-8-sig")
        info = json.loads(raw)
        if isinstance(info, dict) and any("REEMPLAZAR_" in str(value) for value in info.values()):
            raise RuntimeError(
                "Las credenciales todavía contienen la plantilla. "
                "Reemplazala por el JSON de la cuenta de servicio descargado de Google Cloud."
            )
        creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
        _service = build("sheets", "v4", credentials=creds, cache_discovery=False)
    return _service


def read_public_sheet():
    """Lee todas las pestañas mediante la exportación pública, sin credenciales."""
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    try:
        with urlopen(url, timeout=30) as response:
            content = response.read()
    except HTTPError as exc:
        raise RuntimeError(
            f"No se pudo descargar el Sheet (HTTP {exc.code}). "
            "Verificá el SHEET_ID y el acceso 'Cualquier persona con el enlace: Lector'."
        ) from exc
    except (URLError, TimeoutError) as exc:
        raise RuntimeError("No se pudo conectar con Google Sheets. Intentá sincronizar nuevamente.") from exc
    if not content.startswith(b"PK"):
        raise RuntimeError("Google no devolvió una planilla. Verificá el acceso público del Sheet.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        titles, tables = [], []
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = [
                    value.strftime("%d/%m/%Y %H:%M:%S") if isinstance(value, datetime)
                    else value.strftime("%d/%m/%Y") if isinstance(value, date)
                    else "" if value is None else str(value)
                    for value in row
                ]
                while values and values[-1] == "":
                    values.pop()
                rows.append(values)
            while rows and not rows[-1]:
                rows.pop()
            titles.append(worksheet.title)
            tables.append(rows)
        return {"titles": titles, "tables": tables, "fetchedAt": int(time.time() * 1000)}
    finally:
        workbook.close()


def read_sheet():
    """Devuelve todas las hojas del Sheet como lista de tablas (filas de celdas)."""
    if SHEET_ACCESS == "public":
        return read_public_sheet()
    if SHEET_ACCESS != "private":
        raise RuntimeError("SHEET_ACCESS debe ser public o private.")
    svc = sheets_service()
    meta = svc.spreadsheets().get(
        spreadsheetId=SHEET_ID, fields="sheets.properties.title"
    ).execute()
    titles = [s["properties"]["title"] for s in meta.get("sheets", [])]
    ranges = ["'" + t.replace("'", "''") + "'" for t in titles]
    res = svc.spreadsheets().values().batchGet(
        spreadsheetId=SHEET_ID,
        ranges=ranges,
        valueRenderOption="FORMATTED_VALUE",
    ).execute()
    tables = [vr.get("values", []) for vr in res.get("valueRanges", [])]
    return {"titles": titles, "tables": tables, "fetchedAt": int(time.time() * 1000)}


def require_password(view):
    """Protección simple con contraseña (HTTP Basic) si ACCESS_PASSWORD está definida."""
    @wraps(view)
    def wrapper(*args, **kwargs):
        if ACCESS_PASSWORD:
            auth = request.authorization
            if not auth or auth.password != ACCESS_PASSWORD:
                return Response(
                    "Acceso restringido", 401,
                    {"WWW-Authenticate": 'Basic realm="OWDs Del Palacio"'},
                )
        return view(*args, **kwargs)
    return wrapper


@app.route("/")
@require_password
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.route("/api/sheet")
@require_password
def api_sheet():
    force = request.args.get("refresh") == "1"
    with _lock:
        fresh = _cache["data"] and time.time() - _cache["at"] < CACHE_SECONDS
        if force or not fresh:
            try:
                _cache["data"] = read_sheet()
                _cache["at"] = time.time()
            except Exception as exc:  # noqa: BLE001
                app.logger.exception("Error leyendo el Sheet")
                if _cache["data"] is None:
                    return jsonify({"error": str(exc)}), 502
        data = _cache["data"]
    resp = jsonify(data)
    resp.headers["Cache-Control"] = "no-store"
    return resp


@app.route("/health")
def health():
    return {"ok": True}


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=True)
